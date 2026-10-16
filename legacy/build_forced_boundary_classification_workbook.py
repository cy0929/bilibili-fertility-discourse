import argparse
import hashlib
import os
from datetime import datetime
from pathlib import Path

import numpy as np
import pandas as pd
from openpyxl import load_workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter


MODEL = "Qwen/Qwen3-Embedding-8B"
MAX_CHARS_PER_TEXT = 1500

# Embedding cache filename (produced by bertopic_api.py, cached by text hash to avoid duplicate API calls)
EMBEDDING_CACHE_NAME = "embedding_cache_Qwen_Qwen3-Embedding-8B.npz"


def text_hash(text: str, model: str = MODEL) -> str:
    h = hashlib.sha1()
    h.update(model.encode("utf-8"))
    h.update(b"\0")
    h.update(str(text)[:MAX_CHARS_PER_TEXT].encode("utf-8"))
    return h.hexdigest()[:16]


def safe_l2_normalize(x: np.ndarray) -> np.ndarray:
    denom = np.linalg.norm(x, axis=1, keepdims=True)
    denom = np.where(denom == 0, 1.0, denom)
    return x / denom


def first_non_empty(series: pd.Series, default: str = "") -> str:
    for value in series:
        if pd.notna(value) and str(value).strip():
            return str(value)
    return default


def pct(x: int, total: int) -> float:
    return float(x / total) if total else 0.0


def normalize_secondary_no(value) -> str:
    if pd.isna(value):
        return ""
    if isinstance(value, (float, np.floating)):
        return f"{float(value):.1f}"
    return str(value).strip()


def autofit_and_style(path: Path) -> None:
    wb = load_workbook(path)
    header_fill = PatternFill("solid", fgColor="1F4E78")
    header_font = Font(color="FFFFFF", bold=True)
    subheader_fill = PatternFill("solid", fgColor="D9EAF7")

    for ws in wb.worksheets:
        ws.freeze_panes = "A2"
        if ws.max_row >= 1:
            for cell in ws[1]:
                cell.fill = header_fill
                cell.font = header_font
                cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
            ws.auto_filter.ref = ws.dimensions

        for row in ws.iter_rows(min_row=2, max_row=min(ws.max_row, 4)):
            for cell in row:
                cell.alignment = Alignment(vertical="top", wrap_text=True)

        for idx, col_cells in enumerate(ws.columns, start=1):
            header = str(col_cells[0].value or "")
            max_len = len(header)
            for cell in list(col_cells)[1: min(ws.max_row, 80)]:
                value = "" if cell.value is None else str(cell.value)
                max_len = max(max_len, min(len(value), 80))
            width = min(max(max_len + 2, 10), 42)
            if any(key in header for key in ["评论文本", "代表评论", "说明", "解释", "理由"]):
                width = 48
            ws.column_dimensions[get_column_letter(idx)].width = width

        if ws.title == "说明":
            for row in ws.iter_rows(min_row=1, max_row=ws.max_row):
                for cell in row:
                    cell.alignment = Alignment(vertical="top", wrap_text=True)
            for cell in ws["A"]:
                if cell.row > 1:
                    cell.fill = subheader_fill
                    cell.font = Font(bold=True)
            ws.column_dimensions["A"].width = 26
            ws.column_dimensions["B"].width = 88

    wb.save(path)


def _resolve_one(directory: Path, patterns: tuple[str, ...], label: str) -> Path:
    """Find the uniquely matching file under `directory` following the order of `patterns`; when multiple matches exist, take the latest modified one."""
    for pattern in patterns:
        matches = sorted(directory.glob(pattern), key=lambda p: p.stat().st_mtime, reverse=True)
        if matches:
            if len(matches) > 1:
                print(f"[warn] {label} 匹配到 {len(matches)} 个文件，使用最新的：{matches[0].name}")
            return matches[0]
    raise FileNotFoundError(
        f"未找到{label}。在 {directory} 下按 {patterns} 查找均无结果，请用 CLI 参数显式指定。"
    )


def main() -> None:
    parser = argparse.ArgumentParser(
        description="边界样本强制归类工作簿：对 BERTopic 标记为边界/噪声的评论做强制一级二级主题归类"
    )
    parser.add_argument(
        "--input-dir",
        type=Path,
        default=Path(os.environ.get("BFD_TOPIC_OUTPUT_DIR", "output/topic_classification")),
        help="BERTopic 输出目录（环境变量 BFD_TOPIC_OUTPUT_DIR）",
    )
    parser.add_argument("--main-file", type=Path, default=None, help="一级二级主题分类工作簿（含「评论一级二级主题」sheet）")
    parser.add_argument("--boundary-file", type=Path, default=None, help="边界不清样本逐条分析工作簿（含完全不相关标记）")
    parser.add_argument("--cache-file", type=Path, default=None, help=f"嵌入缓存 npz，默认 <input-dir>/{EMBEDDING_CACHE_NAME}")
    parser.add_argument("--out-file", type=Path, default=None, help="输出工作簿路径，默认自动生成带时间戳文件名")
    args = parser.parse_args()

    project_dir = args.input_dir
    output_dir = project_dir
    if not project_dir.is_dir():
        raise NotADirectoryError(
            f"输入目录不存在：{project_dir}\n"
            f"请先运行 bertopic_api.py 产出主题分类结果，或用 --input-dir 指定。"
        )

    main_path = args.main_file or _resolve_one(
        output_dir,
        ("一级二级主题表_*论文命名*.xlsx", "一级二级主题表_*7类*.xlsx", "一级二级主题表_*.xlsx"),
        "一级二级主题分类工作簿",
    )
    boundary_filter_path = args.boundary_file or _resolve_one(
        output_dir,
        ("边界不清样本_完全不相关标记_*.xlsx", "边界不清样本*.xlsx"),
        "边界样本标记工作簿",
    )
    cache_path = args.cache_file or (project_dir / EMBEDDING_CACHE_NAME)
    print(f"[in ] 主题分类工作簿 : {main_path}")
    print(f"[in ] 边界标记工作簿 : {boundary_filter_path}")
    print(f"[in ] 嵌入缓存       : {cache_path}")

    main_sheet = "评论一级二级主题"
    hierarchy_sheet = "一级二级主题表"
    boundary_sheet = "10008条边界样本_带标记"
    filtered_sheet = "完全不相关样本_标1"

    comments = pd.read_excel(main_path, sheet_name=main_sheet)
    old_hierarchy = pd.read_excel(main_path, sheet_name=hierarchy_sheet)
    boundary = pd.read_excel(boundary_filter_path, sheet_name=boundary_sheet)
    filtered_149 = pd.read_excel(boundary_filter_path, sheet_name=filtered_sheet)

    comments["doc_id"] = comments["doc_id"].astype(int)
    comments["二级主题编号"] = comments["二级主题编号"].apply(normalize_secondary_no)
    old_hierarchy["二级主题编号"] = old_hierarchy["二级主题编号"].apply(normalize_secondary_no)
    boundary["评论ID"] = boundary["评论ID"].astype(int)
    filtered_149["评论ID"] = filtered_149["评论ID"].astype(int)

    boundary_ids = set(boundary["评论ID"].tolist())
    filtered_ids = set(boundary.loc[boundary["完全不相关标记"].eq(1), "评论ID"].tolist())
    force_ids = set(boundary.loc[boundary["完全不相关标记"].eq(0), "评论ID"].tolist())

    if len(boundary_ids) != 10008:
        raise RuntimeError(f"边界样本数量不是 10008，而是 {len(boundary_ids)}")
    if len(filtered_ids) != 149:
        raise RuntimeError(f"过滤样本数量不是 149，而是 {len(filtered_ids)}")
    if len(force_ids) != 9859:
        raise RuntimeError(f"强制归类样本数量不是 9859，而是 {len(force_ids)}")
    if not boundary_ids.issubset(set(comments["doc_id"])):
        missing = sorted(boundary_ids - set(comments["doc_id"]))[:10]
        raise RuntimeError(f"边界样本 ID 无法在主表中找到，示例：{missing}")

    boundary_main = comments.loc[comments["doc_id"].isin(boundary_ids), ["doc_id", "text", "一级主题名称"]].copy()
    if not boundary_main["一级主题名称"].eq("边界不清/其他").all():
        bad = boundary_main.loc[~boundary_main["一级主题名称"].eq("边界不清/其他"), "doc_id"].head(10).tolist()
        raise RuntimeError(f"有边界样本在主表中不是“边界不清/其他”，示例：{bad}")

    cache = np.load(cache_path, allow_pickle=False)
    hashes = [text_hash(text) for text in comments["text"].astype(str).tolist()]
    cache_keys = set(cache.files)
    missing_hashes = [h for h in hashes if h not in cache_keys]
    if missing_hashes:
        raise RuntimeError(f"embedding 缓存缺失 {len(missing_hashes)} 条，示例：{missing_hashes[:5]}")

    embeddings = np.vstack([cache[h] for h in hashes]).astype(np.float32)
    vectors = safe_l2_normalize(embeddings)
    doc_to_pos = {int(doc_id): i for i, doc_id in enumerate(comments["doc_id"].tolist())}

    valid_original = comments[
        comments["一级主题名称"].ne("边界不清/其他")
        & comments["二级主题编号"].notna()
        & comments["二级主题编号"].astype(str).str.strip().ne("")
    ].copy()

    sec_info = {}
    centroids_by_primary = {}
    for (primary_no, sec_no), group in valid_original.groupby(["一级主题编号", "二级主题编号"], dropna=False):
        positions = [doc_to_pos[int(x)] for x in group["doc_id"].tolist()]
        centroid = vectors[positions].mean(axis=0, keepdims=True)
        centroid = safe_l2_normalize(centroid)[0]
        primary_no = int(primary_no)
        sec_no = str(sec_no)
        info = {
            "一级主题编号": primary_no,
            "一级主题名称": first_non_empty(group["一级主题名称"]),
            "二级主题编号": sec_no,
            "二级主题名称": first_non_empty(group["二级主题名称"]),
            "自动关键词命名": first_non_empty(group.get("自动关键词命名", pd.Series(dtype=object))),
            "二级主题关键词": first_non_empty(group.get("二级主题关键词", pd.Series(dtype=object))),
            "原始评论数": int(len(group)),
            "centroid": centroid,
        }
        sec_info[(primary_no, sec_no)] = info
        centroids_by_primary.setdefault(primary_no, []).append(info)

    forced_boundary = boundary.loc[boundary["完全不相关标记"].eq(0)].copy()
    forced_boundary["最终一级主题编号"] = forced_boundary["第一候选编号"].astype(int)
    forced_boundary["最终一级主题名称"] = forced_boundary["第一候选一级主题"].astype(str)

    assigned_rows = []
    for _, row in forced_boundary.iterrows():
        doc_id = int(row["评论ID"])
        primary_no = int(row["最终一级主题编号"])
        candidates = centroids_by_primary.get(primary_no, [])
        if not candidates:
            raise RuntimeError(f"一级主题 {primary_no} 没有可用二级主题中心")
        vec = vectors[doc_to_pos[doc_id]]
        centers = np.vstack([c["centroid"] for c in candidates])
        sims = vec @ centers.T
        order = np.argsort(-sims)
        best = candidates[int(order[0])]
        second = candidates[int(order[1])] if len(order) > 1 else None
        second_sim = float(sims[int(order[1])]) if len(order) > 1 else np.nan
        assigned_rows.append({
            "评论ID": doc_id,
            "最终一级主题编号": primary_no,
            "最终一级主题名称": str(row["最终一级主题名称"]),
            "最终二级主题编号": best["二级主题编号"],
            "最终二级主题名称": best["二级主题名称"],
            "最终二级主题关键词": best["二级主题关键词"],
            "二级主题相似度_中心匹配": float(sims[int(order[0])]),
            "第二候选二级主题编号": second["二级主题编号"] if second else "",
            "第二候选二级主题名称": second["二级主题名称"] if second else "",
            "第二候选二级主题相似度": second_sim,
            "二级主题相似度差": float(sims[int(order[0])] - second_sim) if second else np.nan,
        })

    forced_assign = pd.DataFrame(assigned_rows)
    forced_full = forced_boundary.merge(forced_assign, on=["评论ID", "最终一级主题编号", "最终一级主题名称"], how="left")

    updated = comments.loc[~comments["doc_id"].isin(filtered_ids)].copy()
    updated["完全不相关标记"] = 0
    updated["最终纳入分析"] = 1
    updated["分类来源"] = np.where(updated["doc_id"].isin(force_ids), "边界样本强制归类", "原始7类分类")
    updated["原一级主题编号"] = updated["一级主题编号"]
    updated["原一级主题名称"] = updated["一级主题名称"]
    updated["原二级主题编号"] = updated["二级主题编号"]
    updated["原二级主题名称"] = updated["二级主题名称"]
    for col in [
        "一级主题判定原因", "候选第二一级主题", "一级主题强化关键词",
        "二级主题编号", "二级主题名称", "二级主题关键词", "review_reason", "noise_reason",
        "原二级主题编号", "原二级主题名称",
    ]:
        if col in updated.columns:
            updated[col] = updated[col].astype(object)

    forced_by_doc = forced_full.set_index("评论ID")
    for doc_id, row in forced_by_doc.iterrows():
        mask = updated["doc_id"].eq(int(doc_id))
        updated.loc[mask, "一级主题编号"] = int(row["最终一级主题编号"])
        updated.loc[mask, "一级主题名称"] = str(row["最终一级主题名称"])
        updated.loc[mask, "一级主题判定原因"] = "forced_from_boundary_after_filter"
        updated.loc[mask, "一级主题相似度"] = row.get("第一候选原始相似度_重算", np.nan)
        updated.loc[mask, "一级主题边界差距"] = row.get("第一二候选得分差_重算", np.nan)
        updated.loc[mask, "候选第二一级主题"] = row.get("第二候选一级主题_重算", "")
        updated.loc[mask, "一级主题强化关键词"] = row.get("不相关规则命中", "")
        updated.loc[mask, "二级主题编号"] = str(row["最终二级主题编号"])
        updated.loc[mask, "二级主题名称"] = str(row["最终二级主题名称"])
        updated.loc[mask, "二级主题关键词"] = str(row["最终二级主题关键词"])
        updated.loc[mask, "二级主题相似度"] = float(row["二级主题相似度_中心匹配"])
        updated.loc[mask, "needs_manual_review"] = True
        updated.loc[mask, "review_reason"] = "forced_boundary_assignment"
        updated.loc[mask, "noise_reason"] = ""

    updated["一级主题编号"] = updated["一级主题编号"].astype(int)

    forced_detail_cols = [
        "评论ID", "来源文件", "视频ID", "视频标题", "评论类型", "点赞数", "评论文本",
        "完全不相关标记", "最终一级主题编号", "最终一级主题名称", "最终二级主题编号", "最终二级主题名称",
        "最终二级主题关键词", "二级主题相似度_中心匹配", "第二候选二级主题编号", "第二候选二级主题名称",
        "第二候选二级主题相似度", "二级主题相似度差", "上一版边界原因类型", "上一版一级主题相似度",
        "上一版一级主题边界差距", "第一候选原始相似度_重算", "第一候选综合得分_重算",
        "第二候选一级主题_重算", "第二候选原始相似度_重算", "第二候选综合得分_重算",
        "第一二候选得分差_重算", "归类可行性", "人工复核优先级", "强制归类说明",
    ]
    forced_detail = forced_full[[c for c in forced_detail_cols if c in forced_full.columns]].copy()

    analysis_total = len(updated)
    original_counts = valid_original.groupby(["一级主题编号", "一级主题名称"]).size().rename("原始7类评论数").reset_index()
    forced_counts = forced_detail.groupby(["最终一级主题编号", "最终一级主题名称"]).size().rename("新增强制归类数").reset_index()
    forced_counts = forced_counts.rename(columns={"最终一级主题编号": "一级主题编号", "最终一级主题名称": "一级主题名称"})
    primary_summary = (
        updated.groupby(["一级主题编号", "一级主题名称"], dropna=False)
        .agg(
            更新后评论数=("doc_id", "count"),
            平均一级主题相似度=("一级主题相似度", "mean"),
            平均一级边界差距=("一级主题边界差距", "mean"),
            平均二级主题相似度=("二级主题相似度", "mean"),
            需人工复核数=("needs_manual_review", "sum"),
        )
        .reset_index()
        .merge(original_counts, on=["一级主题编号", "一级主题名称"], how="left")
        .merge(forced_counts, on=["一级主题编号", "一级主题名称"], how="left")
    )
    primary_summary["原始7类评论数"] = primary_summary["原始7类评论数"].fillna(0).astype(int)
    primary_summary["新增强制归类数"] = primary_summary["新增强制归类数"].fillna(0).astype(int)
    primary_summary["占更新后分析样本比例"] = primary_summary["更新后评论数"] / analysis_total
    primary_summary = primary_summary[
        [
            "一级主题编号", "一级主题名称", "更新后评论数", "原始7类评论数", "新增强制归类数",
            "占更新后分析样本比例", "平均一级主题相似度", "平均一级边界差距",
            "平均二级主题相似度", "需人工复核数",
        ]
    ].sort_values("一级主题编号")

    original_sec_counts = valid_original.groupby(["一级主题编号", "一级主题名称", "二级主题编号", "二级主题名称"]).size().rename("原始7类评论数").reset_index()
    forced_sec_counts = forced_detail.groupby(["最终一级主题编号", "最终一级主题名称", "最终二级主题编号", "最终二级主题名称"]).size().rename("新增强制归类数").reset_index()
    forced_sec_counts = forced_sec_counts.rename(columns={
        "最终一级主题编号": "一级主题编号",
        "最终一级主题名称": "一级主题名称",
        "最终二级主题编号": "二级主题编号",
        "最终二级主题名称": "二级主题名称",
    })

    sec_summary = (
        updated.groupby(["一级主题编号", "一级主题名称", "二级主题编号", "二级主题名称"], dropna=False)
        .agg(
            更新后评论数=("doc_id", "count"),
            平均二级主题相似度=("二级主题相似度", "mean"),
            平均一级主题相似度=("一级主题相似度", "mean"),
            平均一级边界差距=("一级主题边界差距", "mean"),
        )
        .reset_index()
        .merge(original_sec_counts, on=["一级主题编号", "一级主题名称", "二级主题编号", "二级主题名称"], how="left")
        .merge(forced_sec_counts, on=["一级主题编号", "一级主题名称", "二级主题编号", "二级主题名称"], how="left")
    )
    sec_summary["原始7类评论数"] = sec_summary["原始7类评论数"].fillna(0).astype(int)
    sec_summary["新增强制归类数"] = sec_summary["新增强制归类数"].fillna(0).astype(int)
    primary_totals = sec_summary.groupby("一级主题编号")["更新后评论数"].transform("sum")
    sec_summary["占本一级主题比例"] = sec_summary["更新后评论数"] / primary_totals
    sec_summary["占更新后分析样本比例"] = sec_summary["更新后评论数"] / analysis_total

    old_sec_meta = old_hierarchy[["二级主题编号", "自动关键词命名", "关键词", "代表评论"]].copy()
    old_sec_meta = old_sec_meta.drop_duplicates("二级主题编号")
    sec_summary = sec_summary.merge(old_sec_meta, on="二级主题编号", how="left")
    sec_summary = sec_summary[
        [
            "一级主题编号", "一级主题名称", "二级主题编号", "二级主题名称",
            "更新后评论数", "原始7类评论数", "新增强制归类数",
            "占本一级主题比例", "占更新后分析样本比例",
            "自动关键词命名", "关键词", "代表评论",
            "平均二级主题相似度", "平均一级主题相似度", "平均一级边界差距",
        ]
    ].sort_values(["一级主题编号", "更新后评论数"], ascending=[True, False])

    filtered_detail = filtered_149.copy()
    filtered_detail["过滤后处理"] = "完全不相关/纯玩梗，未纳入更新后主题分析"

    all_status = comments[[
        "doc_id", "source_file", "video_id", "title", "comment_type", "likes", "text",
        "一级主题编号", "一级主题名称", "二级主题编号", "二级主题名称",
        "一级主题相似度", "一级主题边界差距", "二级主题相似度",
    ]].copy()
    all_status["最终状态"] = "原始7类分类"
    all_status.loc[all_status["doc_id"].isin(force_ids), "最终状态"] = "边界样本强制归类"
    all_status.loc[all_status["doc_id"].isin(filtered_ids), "最终状态"] = "已过滤_完全不相关"
    update_lookup = updated.set_index("doc_id")
    for doc_id in force_ids:
        all_status.loc[all_status["doc_id"].eq(doc_id), "一级主题编号"] = update_lookup.loc[doc_id, "一级主题编号"]
        all_status.loc[all_status["doc_id"].eq(doc_id), "一级主题名称"] = update_lookup.loc[doc_id, "一级主题名称"]
        all_status.loc[all_status["doc_id"].eq(doc_id), "二级主题编号"] = update_lookup.loc[doc_id, "二级主题编号"]
        all_status.loc[all_status["doc_id"].eq(doc_id), "二级主题名称"] = update_lookup.loc[doc_id, "二级主题名称"]
        all_status.loc[all_status["doc_id"].eq(doc_id), "一级主题相似度"] = update_lookup.loc[doc_id, "一级主题相似度"]
        all_status.loc[all_status["doc_id"].eq(doc_id), "一级主题边界差距"] = update_lookup.loc[doc_id, "一级主题边界差距"]
        all_status.loc[all_status["doc_id"].eq(doc_id), "二级主题相似度"] = update_lookup.loc[doc_id, "二级主题相似度"]
    all_status.loc[all_status["doc_id"].isin(filtered_ids), ["一级主题编号", "一级主题名称", "二级主题编号", "二级主题名称"]] = [
        -2, "已过滤/完全不相关", "", "",
    ]

    notes = pd.DataFrame([
        ["输出目的", "将原“边界不清/其他”的 10008 条中，去掉 149 条完全不相关/纯玩梗评论；剩余 9859 条强制归入 7 个一级主题，并计算其所属二级主题。"],
        ["一级主题归类", "使用边界逐条分析表中的“第一候选一级主题/若强制归类_建议一级主题”。这些样本原本因相似度不足或第一二候选过近被保留为边界，本文件为后续分析提供强制归类版本。"],
        ["二级主题归类", "在强制归入的一级主题内部，使用已有正式样本计算每个二级主题的归一化 embedding 中心；边界样本归入与其 embedding 余弦相似度最高的二级主题。"],
        ["更新后分析样本", f"{analysis_total} 条，即原正式 7 类样本 {len(valid_original)} 条 + 强制归类样本 {len(force_ids)} 条。"],
        ["过滤样本", f"{len(filtered_ids)} 条，单独放在“过滤149”sheet，不进入更新后主题汇总。"],
        ["总评论数核对", f"{len(comments)} = 更新后分析样本 {analysis_total} + 过滤样本 {len(filtered_ids)}。"],
        ["主要阅读顺序", "先看“一级主题概览_更新”，再看“一级二级主题表_更新”，逐条结果看“评论分类_更新33816”；被过滤样本看“过滤149”。"],
    ], columns=["项目", "说明"])

    timestamp = datetime.now().strftime("%Y%m%d_%H%M%S")
    out_path = args.out_file or output_dir / f"comments_forced_boundary_classification_{MODEL.replace('/', '_')}_{timestamp}.xlsx"
    out_path.parent.mkdir(parents=True, exist_ok=True)

    with pd.ExcelWriter(out_path, engine="openpyxl") as writer:
        notes.to_excel(writer, sheet_name="说明", index=False)
        primary_summary.to_excel(writer, sheet_name="一级主题概览_更新", index=False)
        sec_summary.to_excel(writer, sheet_name="一级二级主题表_更新", index=False)
        updated.to_excel(writer, sheet_name="评论分类_更新33816", index=False)
        forced_detail.to_excel(writer, sheet_name="强制归类9859", index=False)
        filtered_detail.to_excel(writer, sheet_name="过滤149", index=False)
        all_status.to_excel(writer, sheet_name="全部评论状态33965", index=False)

    autofit_and_style(out_path)

    checks = {
        "output": str(out_path),
        "total_comments": int(len(comments)),
        "updated_analysis_rows": int(len(updated)),
        "forced_rows": int(len(forced_detail)),
        "filtered_rows": int(len(filtered_detail)),
        "primary_counts": primary_summary[["一级主题名称", "更新后评论数", "新增强制归类数"]].to_dict("records"),
        "secondary_rows": int(len(sec_summary)),
    }
    print(checks)


if __name__ == "__main__":
    main()
